"""Enhanced Presentation Generator — creates rich, visually appealing PPTs using python-pptx."""

import os
import json
import re
from datetime import datetime
from typing import Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

try:
    import openai
    HAS_OPENAI = True
except ImportError:
    HAS_OPENAI = False


# ─ Color themes for different topics ───────────────────────

THEMES = {
    "tech": {
        "primary": RGBColor(15, 23, 42), "secondary": RGBColor(56, 189, 248), "accent": RGBColor(129, 140, 248),
        "bg_dark": RGBColor(15, 23, 42), "bg_light": RGBColor(30, 41, 59), "text": RGBColor(248, 250, 252),
        "card_bg": RGBColor(51, 65, 85), "name": "Tech & Innovation"
    },
    "business": {
        "primary": RGBColor(30, 58, 95), "secondary": RGBColor(74, 144, 217), "accent": RGBColor(245, 166, 35),
        "bg_dark": RGBColor(30, 58, 95), "bg_light": RGBColor(44, 82, 130), "text": RGBColor(255, 255, 255),
        "card_bg": RGBColor(45, 55, 72), "name": "Business & Strategy"
    },
    "creative": {
        "primary": RGBColor(109, 40, 217), "secondary": RGBColor(167, 139, 250), "accent": RGBColor(244, 114, 182),
        "bg_dark": RGBColor(76, 29, 149), "bg_light": RGBColor(91, 33, 182), "text": RGBColor(250, 245, 255),
        "card_bg": RGBColor(107, 70, 193), "name": "Creative & Design"
    },
    "nature": {
        "primary": RGBColor(6, 95, 70), "secondary": RGBColor(52, 211, 153), "accent": RGBColor(251, 191, 36),
        "bg_dark": RGBColor(6, 78, 59), "bg_light": RGBColor(6, 95, 70), "text": RGBColor(236, 253, 245),
        "card_bg": RGBColor(4, 120, 87), "name": "Nature & Environment"
    },
    "education": {
        "primary": RGBColor(124, 45, 18), "secondary": RGBColor(251, 146, 60), "accent": RGBColor(253, 224, 71),
        "bg_dark": RGBColor(124, 45, 18), "bg_light": RGBColor(154, 52, 18), "text": RGBColor(255, 251, 235),
        "card_bg": RGBColor(146, 64, 14), "name": "Education & Learning"
    },
    "health": {
        "primary": RGBColor(190, 24, 93), "secondary": RGBColor(244, 114, 182), "accent": RGBColor(167, 243, 208),
        "bg_dark": RGBColor(131, 24, 67), "bg_light": RGBColor(157, 23, 77), "text": RGBColor(253, 242, 248),
        "card_bg": RGBColor(190, 24, 93), "name": "Health & Wellness"
    },
    "finance": {
        "primary": RGBColor(30, 39, 97), "secondary": RGBColor(202, 220, 252), "accent": RGBColor(34, 197, 94),
        "bg_dark": RGBColor(30, 39, 97), "bg_light": RGBColor(30, 58, 138), "text": RGBColor(255, 255, 255),
        "card_bg": RGBColor(49, 46, 129), "name": "Finance & Data"
    },
    "general": {
        "primary": RGBColor(51, 65, 85), "secondary": RGBColor(148, 163, 184), "accent": RGBColor(245, 158, 11),
        "bg_dark": RGBColor(30, 41, 59), "bg_light": RGBColor(51, 65, 85), "text": RGBColor(248, 250, 252),
        "card_bg": RGBColor(71, 85, 105), "name": "General"
    },
}


def detect_theme(description: str) -> str:
    """Detect the best theme based on description keywords."""
    desc_lower = description.lower()

    theme_keywords = {
        "tech": ["tech", "software", "ai", "artificial intelligence", "machine learning",
                 "coding", "developer", "app", "digital", "innovation", "llm", "data"],
        "business": ["business", "strategy", "marketing", "sales", "revenue", "growth",
                     "startup", "investor", "pitch", "campaign", "brand"],
        "creative": ["design", "creative", "art", "ui", "ux", "branding", "visual",
                     "aesthetic", "fashion", "media"],
        "nature": ["nature", "environment", "sustainability", "green", "climate",
                   "ecology", "outdoor", "garden"],
        "education": ["education", "learning", "training", "course", "teaching",
                      "student", "academic", "research"],
        "health": ["health", "medical", "wellness", "fitness", "nutrition",
                   "healthcare", "hospital", "pharma"],
        "finance": ["finance", "financial", "banking", "investment", "accounting",
                    "budget", "economic", "money", "fintech"],
    }

    scores = {}
    for theme, keywords in theme_keywords.items():
        scores[theme] = sum(1 for kw in keywords if kw in desc_lower)

    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else "general"


def generate_slide_outline(
    title: str,
    description: str,
    num_slides: int,
    api_key: Optional[str] = None,
) -> list[dict]:
    """Generate a rich slide outline with detailed content."""

    if api_key and HAS_OPENAI:
        return _llm_outline(title, description, num_slides, api_key)

    return _rule_based_outline(title, description, num_slides)


def _llm_outline(title: str, description: str, num_slides: int, api_key: str) -> list[dict]:
    """Use OpenAI to generate rich slide content."""
    try:
        client = openai.OpenAI(api_key=api_key)

        prompt = f"""You are a professional presentation designer. Create a {num_slides}-slide presentation outline.

Title: {title}
Description: {description}

For each slide, provide:
1. A clear slide title (max 8 words)
2. 4-6 detailed bullet points (each 10-20 words, informative and specific)
3. An image search query for finding a relevant visual (5-10 words)

Output ONLY valid JSON in this format:
[
  {{
    "title": "Slide Title",
    "content": [
      "First detailed bullet point with specific information",
      "Second bullet point with data or examples",
      "Third bullet point with actionable insight",
      "Fourth bullet point with supporting detail"
    ],
    "image_query": "relevant search terms for image"
  }}
]

Rules:
- First slide should be a cover/title slide with 1-2 subtitle bullets
- Last slide should be a closing/thank you slide
- Middle slides should have 4-6 substantial bullet points each
- Make content informative, specific, and valuable
- Avoid generic statements - use concrete examples and data
- Image queries should be specific and visual"""

        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            max_tokens=2000,
            temperature=0.7,
        )

        content = response.choices[0].message.content.strip()

        # Extract JSON from response
        json_match = re.search(r'\[.*\]', content, re.DOTALL)
        if json_match:
            slides = json.loads(json_match.group())
            return slides

        return _rule_based_outline(title, description, num_slides)

    except Exception as e:
        print(f"LLM outline failed: {e}")
        return _rule_based_outline(title, description, num_slides)


def _rule_based_outline(title: str, description: str, num_slides: int) -> list[dict]:
    """Generate rich outline using templates."""
    slides = []

    # Cover slide
    slides.append({
        "title": title,
        "content": [
            description[:150] if len(description) > 150 else description,
            "Comprehensive overview and analysis",
            "Key insights and actionable recommendations",
        ],
        "image_query": f"{title} professional",
        "is_cover": True,
    })

    # Content slides
    content_slides = num_slides - 2

    slide_templates = [
        {
            "title": "Overview & Context",
            "bullets": [
                f"Current landscape of {title.lower()}",
                "Key drivers and market forces at play",
                "Historical development and evolution",
                "Stakeholders and ecosystem participants",
                "Current challenges and opportunities",
            ]
        },
        {
            "title": "Key Components",
            "bullets": [
                "Core elements and fundamental building blocks",
                "Primary technologies and methodologies",
                "Essential resources and infrastructure",
                "Human capital and expertise required",
                "Supporting systems and frameworks",
            ]
        },
        {
            "title": "Benefits & Value Proposition",
            "bullets": [
                "Direct benefits and immediate impact",
                "Long-term strategic advantages",
                "Cost savings and efficiency gains",
                "Competitive differentiation",
                "Measurable outcomes and ROI",
            ]
        },
        {
            "title": "Implementation Strategy",
            "bullets": [
                "Phased rollout approach",
                "Key milestones and timelines",
                "Resource allocation and budgeting",
                "Risk mitigation strategies",
                "Success metrics and KPIs",
            ]
        },
        {
            "title": "Challenges & Solutions",
            "bullets": [
                "Common obstacles and pain points",
                "Proven solutions and best practices",
                "Lessons learned from industry leaders",
                "Adaptive strategies for changing conditions",
                "Continuous improvement frameworks",
            ]
        },
        {
            "title": "Future Outlook",
            "bullets": [
                "Emerging trends and innovations",
                "Technology evolution and adoption",
                "Market trajectory and projections",
                "Opportunities for early adopters",
                "Strategic recommendations for stakeholders",
            ]
        },
    ]

    for i in range(content_slides):
        template = slide_templates[i % len(slide_templates)]
        slides.append({
            "title": template["title"],
            "content": template["bullets"],
            "image_query": f"{title} {template['title'].lower()}",
        })

    # Closing slide
    slides.append({
        "title": "Thank You",
        "content": [
            "Questions & Discussion",
            "Contact information",
            "Next steps and resources",
        ],
        "image_query": "thank you professional",
        "is_closing": True,
    })

    return slides


def _add_background(slide, color):
    """Set slide background color."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text_box(slide, left, top, width, height, text, font_size=18, bold=False, color=None, alignment=PP_ALIGN.LEFT):
    """Add a text box to the slide."""
    textbox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    text_frame = textbox.text_frame
    text_frame.word_wrap = True

    p = text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    if color:
        p.font.color.rgb = color
    p.alignment = alignment

    return textbox


def _build_cover_slide(prs, slide_idx, slide_data, theme):
    """Build a cover/title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    _add_background(slide, theme["bg_dark"])

    # Title
    _add_text_box(
        slide, 1, 2.5, 8, 1.5,
        slide_data["title"],
        font_size=44, bold=True, color=theme["text"], alignment=PP_ALIGN.CENTER
    )

    # Subtitle/content
    if slide_data.get("content"):
        subtitle = "\n".join(slide_data["content"][:2])
        _add_text_box(
            slide, 1.5, 4.5, 7, 1,
            subtitle,
            font_size=20, color=theme["secondary"], alignment=PP_ALIGN.CENTER
        )


def _build_content_slide(prs, slide_idx, slide_data, theme, slide_num):
    """Build a content slide with title and bullets."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    _add_background(slide, theme["bg_light"])

    # Title
    _add_text_box(
        slide, 0.5, 0.5, 9, 0.8,
        slide_data["title"],
        font_size=32, bold=True, color=theme["text"], alignment=PP_ALIGN.LEFT
    )

    # Content bullets
    if slide_data.get("content"):
        top = 1.5
        for bullet in slide_data["content"]:
            # Add bullet point
            _add_text_box(
                slide, 0.8, top, 8.4, 0.5,
                f"• {bullet}",
                font_size=16, color=theme["text"], alignment=PP_ALIGN.LEFT
            )
            top += 0.6


def _build_closing_slide(prs, slide_idx, slide_data, theme):
    """Build a closing/thank you slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    _add_background(slide, theme["bg_dark"])

    # Title
    _add_text_box(
        slide, 1, 2.5, 8, 1.5,
        slide_data["title"],
        font_size=48, bold=True, color=theme["text"], alignment=PP_ALIGN.CENTER
    )

    # Content
    if slide_data.get("content"):
        content = "\n".join(slide_data["content"])
        _add_text_box(
            slide, 1.5, 4.5, 7, 1,
            content,
            font_size=18, color=theme["secondary"], alignment=PP_ALIGN.CENTER
        )


def generate_presentation(
    title: str,
    description: str,
    num_slides: int = 6,
    output_dir: Optional[str] = None,
    api_key: Optional[str] = None,
    image_paths: Optional[dict] = None,
) -> str:
    """Generate a rich, visually appealing presentation PPT."""

    if output_dir is None:
        output_dir = os.path.expanduser("~/Desktop")

    # Detect theme
    theme_key = detect_theme(description)
    theme = THEMES[theme_key]

    # Generate slide outline
    print(f"Generating outline for: {title}")
    slides_data = generate_slide_outline(title, description, num_slides, api_key)

    # Create filename
    safe_title = re.sub(r'[^\w\s-]', '', title)[:50]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"{safe_title}_{timestamp}.pptx"
    filepath = os.path.join(output_dir, filename)

    # Build PPT using python-pptx
    print(f"Building presentation with '{theme['name']}' theme...")
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    for i, slide_data in enumerate(slides_data):
        is_cover = slide_data.get("is_cover", False)
        is_closing = slide_data.get("is_closing", False)

        if is_cover:
            _build_cover_slide(prs, i + 1, slide_data, theme)
        elif is_closing:
            _build_closing_slide(prs, i + 1, slide_data, theme)
        else:
            _build_content_slide(prs, i + 1, slide_data, theme, i)

    # Save presentation
    prs.save(filepath)
    print(f"Presentation saved: {filepath}")

    return filepath


def get_available_themes() -> list[dict]:
    """Get list of available themes."""
    return [
        {"key": key, "name": theme["name"]}
        for key, theme in THEMES.items()
    ]
